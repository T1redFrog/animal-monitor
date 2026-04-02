import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import os
from datetime import datetime

def generate_excel_report(history):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт"

    # Title
    ws.merge_cells('A1:F1')
    ws['A1'] = 'Отчёт: Мониторинг бездомных животных во дворах'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:F2')
    ws['A2'] = f'Сформирован: {datetime.now().strftime("%d.%m.%Y %H:%M")}'
    ws['A2'].alignment = Alignment(horizontal='center')

    # Header
    headers = ['№', 'Дата/Время', 'Файл', 'Тип', 'Кошек', 'Собак', 'Всего']
    ws.append([])
    ws.append(headers)

    header_row = 4
    header_fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, _ in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = border

    # Data
    total_cats = total_dogs = total_all = 0
    for i, record in enumerate(history, 1):
        stats = record.get('stats', {})
        cats = stats.get('Cat', 0)
        dogs = stats.get('Dog', 0)
        total = stats.get('total', 0)
        total_cats += cats
        total_dogs += dogs
        total_all += total

        row_data = [
            i,
            record.get('timestamp', ''),
            record.get('filename', ''),
            'Видео' if record.get('type') == 'video' else 'Изображение',
            cats, dogs, total
        ]
        ws.append(row_data)
        row_idx = ws.max_row
        fill = PatternFill(start_color='EBF3FB', end_color='EBF3FB', fill_type='solid') if i % 2 == 0 else None
        for col in range(1, 8):
            cell = ws.cell(row=row_idx, column=col)
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
            if fill:
                cell.fill = fill

    # Summary row
    ws.append([])
    summary_row = ws.max_row + 1
    ws.cell(row=summary_row, column=1, value='ИТОГО')
    ws.cell(row=summary_row, column=5, value=total_cats)
    ws.cell(row=summary_row, column=6, value=total_dogs)
    ws.cell(row=summary_row, column=7, value=total_all)
    for col in range(1, 8):
        cell = ws.cell(row=summary_row, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='D6E4F0', end_color='D6E4F0', fill_type='solid')
        cell.border = border
        cell.alignment = Alignment(horizontal='center')

    # Column widths
    col_widths = [5, 20, 30, 14, 10, 10, 10]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    path = 'report.xlsx'
    wb.save(path)
    return path


def generate_pdf_report(history):
    path = 'report.pdf'
    doc = SimpleDocTemplate(path, pagesize=A4,
                            rightMargin=2*cm, leftMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)

    # Try to register a Cyrillic font
    font_name = 'Helvetica'
    try:
        font_paths = [
            '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
            '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
            '/System/Library/Fonts/Helvetica.ttc',
            'C:/Windows/Fonts/arial.ttf',
        ]
        for fp in font_paths:
            if os.path.exists(fp):
                pdfmetrics.registerFont(TTFont('CyrillicFont', fp))
                font_name = 'CyrillicFont'
                break
    except:
        pass

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('title', fontName=font_name, fontSize=14,
                                  spaceAfter=6, alignment=1, leading=18)
    normal_style = ParagraphStyle('normal', fontName=font_name, fontSize=10,
                                   spaceAfter=4, leading=14)
    header_style = ParagraphStyle('header', fontName=font_name, fontSize=11,
                                   spaceBefore=10, spaceAfter=6, leading=14)

    elements = []
    elements.append(Paragraph('Отчёт: Мониторинг бездомных животных во дворах', title_style))
    elements.append(Paragraph(f'Сформирован: {datetime.now().strftime("%d.%m.%Y %H:%M")}', normal_style))
    elements.append(Spacer(1, 0.5*cm))

    # Stats summary
    total_cats = sum(r.get('stats', {}).get('Cat', 0) for r in history)
    total_dogs = sum(r.get('stats', {}).get('Dog', 0) for r in history)
    elements.append(Paragraph(f'Всего запросов: {len(history)}', normal_style))
    elements.append(Paragraph(f'Обнаружено кошек: {total_cats}', normal_style))
    elements.append(Paragraph(f'Обнаружено собак: {total_dogs}', normal_style))
    elements.append(Spacer(1, 0.5*cm))

    # Table
    elements.append(Paragraph('История запросов:', header_style))
    table_data = [['№', 'Дата/Время', 'Файл', 'Тип', 'Кошек', 'Собак', 'Всего']]

    for i, record in enumerate(history, 1):
        stats = record.get('stats', {})
        table_data.append([
            str(i),
            record.get('timestamp', ''),
            record.get('filename', '')[:30],
            'Видео' if record.get('type') == 'video' else 'Фото',
            str(stats.get('Cat', 0)),
            str(stats.get('Dog', 0)),
            str(stats.get('total', 0)),
        ])

    t = Table(table_data, colWidths=[1*cm, 4*cm, 5*cm, 2.5*cm, 1.8*cm, 1.8*cm, 1.8*cm])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2E75B6')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EBF3FB')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWHEIGHT', (0, 0), (-1, -1), 18),
    ]))
    elements.append(t)

    doc.build(elements)
    return path
